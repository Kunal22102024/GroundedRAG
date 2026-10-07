"""Document parsers: PDF, DOCX, Markdown, HTML and plain text -> page-aware text segments.

Every parser returns a ``ParsedDocument``: an ordered list of ``Segment`` (one
paragraph, list item, table row or code block, with the 1-based physical page it
sits on and the heading it falls under) plus the file's own properties, which the
metadata resolver uses as a fallback.

Error messages on ``ParseError`` are written for the end user and never contain
document text (CLAUDE.md rule 21).
"""

from __future__ import annotations

import logging
import math
import re
import unicodedata
import zipfile
from collections import Counter
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path
from statistics import median
from typing import Any

from app.constants import (
    MIN_PAGE_TEXT_CHARS,
    SCANNED_PDF_MAX_TEXT_PAGE_RATIO,
    SCANNED_PDF_MESSAGE,
)

logger = logging.getLogger(__name__)

# pdfminer warns once per font with a missing bounding box ("Could not get FontBBox"):
# hundreds of lines for a large manual and nothing anyone can act on.
for _noisy in ("pdfminer", "pdfplumber"):
    logging.getLogger(_noisy).setLevel(logging.ERROR)


class ParseError(Exception):
    """A document could not be parsed. The message is safe to show to the user."""


class ScannedPdfError(ParseError):
    """A PDF has no (or almost no) extractable text; OCR is not supported."""


class EmptyDocumentError(ParseError):
    """The document parsed but contains no text."""


class UnsupportedFormatError(ParseError):
    """The file extension is not one of the supported formats."""


@dataclass(frozen=True)
class Segment:
    """One block of body text in reading order."""

    page: int  # 1-based physical page; always 1 for formats without pages
    section: str  # nearest heading above this text, or ""
    text: str


@dataclass(frozen=True)
class DocumentProperties:
    """What the file says about itself (metadata tier 3, after upload values and file name)."""

    title: str = ""
    created: date | None = None
    department: str = ""
    document_type: str = ""


@dataclass(frozen=True)
class ParsedDocument:
    """Parser output."""

    segments: list[Segment]
    properties: DocumentProperties
    page_count: int = 1
    # Diagnostics (counts only, never text): e.g. {"outline_entries": 808, "headings_matched": 790}.
    stats: dict[str, int] = field(default_factory=dict)


# --- Text helpers --------------------------------------------------------------

_LIGATURES = str.maketrans({"ﬀ": "ff", "ﬁ": "fi", "ﬂ": "fl", "ﬃ": "ffi", "ﬄ": "ffl"})
_INVISIBLE = re.compile("[­​‌‍⁠﻿]")
_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")
_INLINE_SPACE = re.compile(r"[ \t  -   　]+")
_BULLET = re.compile(r"^(?:[•◦▪■●○▫‣⁃*\-–—]|\d{1,3}[.)]|[a-zA-Z][.)]|\(\d{1,3}\))\s+")
_SENTENCE_END = re.compile(r"[.!?:;…][\"')\]]*$")


def clean_text(text: str) -> str:
    """Normalize one block: ligatures, invisible/control characters, runs of spaces."""
    text = unicodedata.normalize("NFC", text).translate(_LIGATURES)
    text = _INVISIBLE.sub("", _CONTROL.sub(" ", text))
    text = _INLINE_SPACE.sub(" ", text)
    return "\n".join(line.strip() for line in text.split("\n")).strip()


def _join_wrapped(lines: list[str]) -> str:
    """Join hard-wrapped lines into one paragraph, repairing end-of-line hyphenation."""
    out = ""
    for line in lines:
        line = line.strip()
        if not line:
            continue
        if not out:
            out = line
        elif out.endswith("-") and len(out) > 1 and out[-2].isalpha() and line[:1].islower():
            out = out[:-1] + line  # "exam-" + "ple" -> "example"
        else:
            out = f"{out} {line}"
    return out


def _letters(text: str) -> int:
    return sum(ch.isalpha() for ch in text)


def _normalize_title(text: str) -> str:
    """Case/space/punctuation-insensitive key for matching outline titles to page lines."""
    return re.sub(r"[\W_]+", " ", text.casefold()).strip()


# --- PDF -------------------------------------------------------------------------

_BOLD_FONT = re.compile(r"bold|black|heavy|semibold|demi|-bd\b|-bol\b", re.IGNORECASE)
_PDF_DATE = re.compile(r"D:(\d{4})(\d{2})?(\d{2})?")
_MAX_HEADING_CHARS = 120
_HEADING_SIZE_RATIO = 1.12
_MIN_GUTTER_PT = 10.0  # narrowest empty vertical strip that can separate two text columns
_MAX_SPANNING_SHARE = 0.10  # lines allowed to cross the gutter (full-width headings)
_MIN_COLUMN_LINE_CHARS = 20  # typical line length per column; shorter means table cells, not prose
_PAGE_LABEL = re.compile(
    r"^(?:(?:page|pg\.?|p\.)\s*\d+(?:\s*(?:of|/)\s*\d+)?|[-\u2013\u2014\s]*\d{1,4}[-\u2013\u2014\s]*|\d{1,4}\s*(?:of|/)\s*\d{1,4})$",
    re.IGNORECASE,
)
_MAX_CHROME_SHARE = 0.3  # never strip more than this share of a PDF's text as 'page chrome'


@dataclass
class _Line:
    text: str
    size: float
    bold: bool
    top: float
    bottom: float
    heading: bool = False


@dataclass(frozen=True)
class _OutlineEntry:
    level: int
    title: str
    page: int


def _parse_pdf_date(value: object) -> date | None:
    match = _PDF_DATE.match(str(value or ""))
    if not match:
        return None
    year, month, day = int(match.group(1)), int(match.group(2) or 1), int(match.group(3) or 1)
    try:
        return date(year, month, day)
    except ValueError:
        return None


def _line_from_chars(chars: list[dict[str, Any]], text: str | None = None) -> _Line | None:
    """Build a line from pdfplumber characters (all on one visual line)."""
    from pdfplumber.utils import extract_text

    content = clean_text(text if text is not None else extract_text(chars))
    if not content:
        return None
    sizes = [float(c["size"]) for c in chars if c.get("size")]
    size = round(float(median(sizes)), 1) if sizes else 0.0
    bold_chars = sum(1 for c in chars if _BOLD_FONT.search(c.get("fontname", "") or ""))
    bold = bool(chars) and bold_chars / len(chars) >= 0.8
    return _Line(
        content,
        size,
        bold,
        min(float(c["top"]) for c in chars),
        max(float(c["bottom"]) for c in chars),
    )


def _find_gutter(raw_lines: list[dict[str, Any]], page_width: float) -> float | None:
    """The x position of the gutter between two text columns, or ``None`` for one column.

    A gutter is a vertical strip at least ``_MIN_GUTTER_PT`` wide that almost no line crosses
    (a few full-width headings may), with substantial prose on both sides. Pages whose "columns"
    hold short cells (tables, forms) are left alone: splitting those would scramble rows.
    """
    import numpy as np

    lines = [[c for c in raw["chars"] if c["text"].strip()] for raw in raw_lines]
    lines = [chars for chars in lines if chars]
    if len(lines) < 6:
        return None
    width = int(page_width) + 2
    cover = np.zeros(width, dtype=np.int32)
    for chars in lines:
        row = np.zeros(width, dtype=bool)
        for c in chars:
            row[int(c["x0"]) : int(c["x1"]) + 1] = True
        cover += row
    low = min(c["x0"] for chars in lines for c in chars)
    high = max(c["x1"] for chars in lines for c in chars)
    lo, hi = int(low + 0.2 * (high - low)), int(low + 0.8 * (high - low))
    allowed = max(1, int(_MAX_SPANNING_SHARE * len(lines)))

    best: tuple[int, int] | None = None
    run_start: int | None = None
    for x in range(lo, hi + 2):
        empty = x <= hi and cover[x] <= allowed
        if empty and run_start is None:
            run_start = x
        elif not empty and run_start is not None:
            if x - run_start >= _MIN_GUTTER_PT and (
                best is None or x - run_start > best[1] - best[0]
            ):
                best = (run_start, x)
            run_start = None
    if best is None:
        return None
    mid = (best[0] + best[1]) / 2

    left_lengths: list[int] = []
    right_lengths: list[int] = []
    for chars in lines:
        left = sum(1 for c in chars if (c["x0"] + c["x1"]) / 2 < mid)
        right = len(chars) - left
        if left:
            left_lengths.append(left)
        if right:
            right_lengths.append(right)
    total = sum(len(chars) for chars in lines)
    if min(sum(left_lengths), sum(right_lengths)) < 0.25 * total:
        return None
    if min(median(left_lengths), median(right_lengths)) < _MIN_COLUMN_LINE_CHARS:
        return None
    return mid


def _columns_in_reading_order(raw_lines: list[dict[str, Any]], mid: float) -> list[_Line]:
    """Left column, then right column, between full-width lines (which separate blocks)."""
    out: list[_Line] = []
    left: list[_Line] = []
    right: list[_Line] = []

    def flush() -> None:
        out.extend(left)
        out.extend(right)
        left.clear()
        right.clear()

    for raw in raw_lines:
        chars = raw["chars"]
        solid = [c for c in chars if c["text"].strip()]
        if not solid:
            continue
        if any(c["x0"] < mid - 1 and c["x1"] > mid + 1 for c in solid):
            flush()  # crosses the gutter: a heading or caption spanning both columns
            if line := _line_from_chars(chars, raw.get("text")):
                out.append(line)
            continue
        on_left = [c for c in chars if (c["x0"] + c["x1"]) / 2 < mid]
        on_right = [c for c in chars if (c["x0"] + c["x1"]) / 2 >= mid]
        if on_left and (line := _line_from_chars(on_left)):
            left.append(line)
        if on_right and (line := _line_from_chars(on_right)):
            right.append(line)
    flush()
    return out


def _extract_page(page: Any) -> tuple[list[_Line], bool]:
    """Text lines of one page in reading order, and whether it was read as two columns."""
    raw_lines = page.extract_text_lines(strip=True, return_chars=True)
    mid = _find_gutter(raw_lines, float(page.width))
    if mid is not None:
        return _columns_in_reading_order(raw_lines, mid), True
    lines = [_line_from_chars(raw.get("chars") or [], raw.get("text", "")) for raw in raw_lines]
    return [line for line in lines if line is not None], False


def _extract_lines(page: Any) -> list[_Line]:
    """Text lines of one page with font size and boldness, in reading order."""
    return _extract_page(page)[0]


def _read_outline(pdf: Any) -> list[_OutlineEntry]:
    """PDF bookmarks resolved to 1-based page numbers (empty if the PDF has none)."""
    try:
        raw_entries = list(pdf.doc.get_outlines())
    except Exception:  # PDFNoOutlines, or a malformed outline tree
        return []
    page_ids = {page.page_obj.pageid: index + 1 for index, page in enumerate(pdf.pages)}
    entries: list[_OutlineEntry] = []
    for level, title, dest, action, _ in raw_entries:
        title = clean_text(str(title or ""))
        try:
            target = dest
            if target is None and action is not None:
                resolved = action.resolve() if hasattr(action, "resolve") else action
                target = resolved.get("D") if isinstance(resolved, dict) else None
            if isinstance(target, (bytes, str)):
                target = pdf.doc.get_dest(target)
            if hasattr(target, "resolve"):
                target = target.resolve()
            if isinstance(target, dict):
                target = target.get("D")
            page = page_ids.get(target[0].objid) if target else None
        except Exception:
            page = None
        if title and page:
            entries.append(_OutlineEntry(int(level), title, page))
    return entries


def _strip_repeated_edges(pages: list[list[_Line]]) -> int:
    """Drop page chrome: running headers/footers and banners that repeat on most pages.

    Digits are normalized so "Page 3" and "Page 4" count as the same line. Top and bottom
    lines are removed when they repeat on at least 40% of pages. Longer lines (12+ characters)
    are removed wherever they sit, because pages printed from a web page carry banners in
    the middle of the page. Needs at least four pages to be sure. Returns how many lines
    were removed.
    """
    labels = 0
    for index, lines in enumerate(pages):
        edge_ids = {id(line) for line in (lines[:2] + lines[-2:])}
        kept = [
            line
            for line in lines
            if id(line) not in edge_ids or not _PAGE_LABEL.match(line.text.strip())
        ]
        labels += len(lines) - len(kept)
        pages[index] = kept
    if len(pages) < 4:
        return labels

    def key(line: _Line) -> str:
        return re.sub(r"\d+", "#", line.text.casefold()).strip()

    def edges(lines: list[_Line]) -> list[_Line]:
        return lines[:2] + lines[-2:] if len(lines) > 4 else list(lines)

    edge_seen: Counter[str] = Counter()
    anywhere_seen: Counter[str] = Counter()
    for lines in pages:
        edge_seen.update({key(line) for line in edges(lines)})
        anywhere_seen.update({key(line) for line in lines})
    threshold = max(3, math.ceil(0.4 * len(pages)))
    repeated_edges = {
        k for k, n in edge_seen.items() if n >= threshold and _letters(k) + k.count("#") > 0
    }
    repeated_anywhere = {
        k for k, n in anywhere_seen.items() if n >= threshold and len(k) >= 12 and _letters(k) >= 8
    }

    plan: list[set[int]] = []
    removed_chars = total_chars = 0
    for lines in pages:
        at_edge = {id(line) for line in edges(lines)}
        drop = {
            id(line)
            for line in lines
            if key(line) in repeated_anywhere
            or (id(line) in at_edge and key(line) in repeated_edges)
        }
        plan.append(drop)
        total_chars += sum(len(line.text) for line in lines)
        removed_chars += sum(len(line.text) for line in lines if id(line) in drop)

    # Real page chrome is a small fraction of a document. If a rule wants to delete a large
    # share (identical pages, forms, tables of near-identical rows), it has mistaken body
    # text for chrome, so remove nothing.
    if total_chars == 0 or removed_chars > _MAX_CHROME_SHARE * total_chars:
        return labels

    removed = 0
    for index, drop in enumerate(plan):
        removed += len(drop)
        pages[index] = [line for line in pages[index] if id(line) not in drop]
    return removed + labels


def _body_size(pages: list[list[_Line]]) -> float:
    """The font size most body text uses (character-weighted mode)."""
    weights: Counter[float] = Counter()
    for lines in pages:
        for line in lines:
            if line.size:
                weights[round(line.size * 2) / 2] += len(line.text)
    return weights.most_common(1)[0][0] if weights else 0.0


def _looks_like_heading(line: _Line, body: float) -> bool:
    text = line.text
    if len(text) > _MAX_HEADING_CHARS or _letters(text) < 2 or _BULLET.match(text):
        return False
    if text.endswith((",", ";")):
        return False
    larger = body > 0 and line.size >= body * _HEADING_SIZE_RATIO
    bold_line = line.bold and len(text) <= 90 and not text.endswith((".", ":"))
    return larger or bold_line


def _parse_pdf(path: Path) -> ParsedDocument:
    import pdfplumber
    from pdfminer.pdfdocument import PDFPasswordIncorrect

    try:
        pdf = pdfplumber.open(path)
    except PDFPasswordIncorrect as exc:
        raise ParseError("The PDF is password-protected and cannot be read.") from exc
    except Exception as exc:
        raise ParseError("The file could not be read as a PDF (it may be corrupt).") from exc

    with pdf:
        page_count = len(pdf.pages)
        if page_count == 0:
            raise EmptyDocumentError("The PDF contains no pages.")

        pages: list[list[_Line]] = []
        failed_pages = 0
        two_column_pages = 0
        for index, page in enumerate(pdf.pages):
            try:
                lines, columns = _extract_page(page)
                pages.append(lines)
                two_column_pages += columns
            except Exception as exc:
                failed_pages += 1
                logger.warning(
                    "pdf page extraction failed",
                    extra={"page": index + 1, "error": type(exc).__name__},
                )
                pages.append([])
            finally:
                page.flush_cache()

        outline = _read_outline(pdf)
        metadata = pdf.metadata or {}

    # Scanned (image-only) detection happens before any cleanup, on what the PDF really holds.
    text_pages = sum(
        1
        for lines in pages
        if sum(len(ln.text.replace(" ", "")) for ln in lines) >= MIN_PAGE_TEXT_CHARS
    )
    if text_pages / page_count < SCANNED_PDF_MAX_TEXT_PAGE_RATIO:
        raise ScannedPdfError(SCANNED_PDF_MESSAGE)

    total_chars = sum(len(ln.text) for lines in pages for ln in lines)
    cid_chars = sum(ln.text.count("(cid:") * 8 for lines in pages for ln in lines)
    if total_chars and cid_chars / total_chars > 0.3:
        raise ParseError("The PDF's text uses a font encoding that cannot be extracted.")

    stats: dict[str, int] = {
        "pages_without_text": page_count - text_pages,
        "failed_pages": failed_pages,
        "two_column_pages": two_column_pages,
    }
    stats["header_footer_lines_removed"] = _strip_repeated_edges(pages)

    use_outline = len(outline) >= max(3, page_count // 10)
    stats["outline_entries"] = len(outline) if use_outline else 0
    body = _body_size(pages)
    if not use_outline:
        for lines in pages:
            for line in lines:
                line.heading = _looks_like_heading(line, body)

    segments, headings_found, matched = _assemble_pdf_segments(
        pages, outline if use_outline else [], body
    )
    stats["headings"] = headings_found
    if use_outline:
        stats["outline_entries_matched"] = matched

    if not segments:
        raise ScannedPdfError(SCANNED_PDF_MESSAGE)

    title = clean_text(str(metadata.get("Title") or ""))
    if not title and segments:
        title = next((seg.section for seg in segments if seg.section), "")
    properties = DocumentProperties(
        title=title[:300],
        created=_parse_pdf_date(metadata.get("CreationDate"))
        or _parse_pdf_date(metadata.get("ModDate")),
    )
    return ParsedDocument(segments, properties, page_count, stats)


class _PdfAssembler:
    """Turns per-page lines into paragraph segments labelled with their heading.

    Headings come from the PDF outline when it has enough entries (matched to the page
    line carrying the same text), otherwise from font size and boldness (``line.heading``).
    """

    def __init__(self, outline: list[_OutlineEntry], body: float) -> None:
        self.outline = bool(outline)
        self.body = body
        self.by_page: dict[int, list[_OutlineEntry]] = {}
        for entry in outline:
            self.by_page.setdefault(entry.page, []).append(entry)
        self.segments: list[Segment] = []
        self.section = ""
        self.headings = 0
        self.matched = 0
        self._page = 1
        self._paragraph: list[_Line] = []
        self._heading_run: list[_Line] = []

    def _flush_paragraph(self) -> None:
        if self._paragraph:
            text = clean_text(_join_wrapped([line.text for line in self._paragraph]))
            if text:
                self.segments.append(Segment(self._page, self.section, text))
            self._paragraph.clear()

    def _close_heading_run(self) -> None:
        if self._heading_run:
            self.section = clean_text(" ".join(line.text for line in self._heading_run))[:300]
            self.headings += 1
            self._heading_run.clear()

    def _starts_new_paragraph(
        self, line: _Line, previous: _Line | None, longest: int, gap_limit: float
    ) -> bool:
        if previous is None or _BULLET.match(line.text):
            return True
        if line.top < previous.top - 1:
            return True  # reading order jumped back up the page: a new column
        if line.top - previous.bottom > gap_limit:
            return True  # a visible vertical gap
        # A short line that ended a sentence, followed by a capital: the paragraph ended.
        return (
            _SENTENCE_END.search(previous.text) is not None
            and line.text[:1].isupper()
            and len(previous.text) < 0.6 * longest
        )

    @staticmethod
    def _match_outline(
        pending: list[_OutlineEntry], lines: list[_Line], start: int
    ) -> tuple[int, int] | None:
        """Find an outline entry whose title is the text of 1 to 3 consecutive lines starting here.

        Titles wrap over lines in print layouts and sometimes lose their spaces, so the
        comparison ignores whitespace and punctuation. Returns (pending index, lines used).
        """
        compact_titles = [_normalize_title(entry.title).replace(" ", "") for entry in pending]
        joined = ""
        for used in range(1, 4):
            if start + used > len(lines):
                break
            joined += _normalize_title(lines[start + used - 1].text).replace(" ", "")
            if joined in compact_titles:
                return compact_titles.index(joined), used
        return None

    def add_page(self, page_number: int, lines: list[_Line]) -> None:
        self._page = page_number
        pending = list(self.by_page.get(page_number, []))
        longest = max((len(line.text) for line in lines), default=0)
        # Line spacing differs between documents, so a paragraph gap is judged against this
        # page's own typical gap between consecutive lines, not a fixed size.
        gaps = [
            b.top - a.bottom for a, b in zip(lines, lines[1:], strict=False) if b.top - a.bottom > 0
        ]
        typical = median(gaps) if len(gaps) >= 3 else 0.0
        gap_limit = max(0.35 * self.body, 1.6 * typical) if typical else 0.8 * self.body
        previous: _Line | None = None

        skip = 0
        for position, line in enumerate(lines):
            if skip:
                skip -= 1
                continue
            if self.outline:
                found = self._match_outline(pending, lines, position) if pending else None
                if found is not None:
                    hit, used = found
                    self._flush_paragraph()
                    self.section = pending[hit].title[:300]
                    del pending[: hit + 1]  # outline entries before the match were passed over
                    self.matched += 1
                    self.headings += 1
                    previous = None
                    skip = used - 1
                    continue
                is_heading = False
            else:
                is_heading = line.heading

            if is_heading:
                self._flush_paragraph()
                if self._heading_run and abs(line.size - self._heading_run[-1].size) > 0.6:
                    self._close_heading_run()  # a different style starts a new heading
                self._heading_run.append(line)
                previous = None
                continue

            self._close_heading_run()
            if self._starts_new_paragraph(line, previous, longest, gap_limit):
                self._flush_paragraph()
            self._paragraph.append(line)
            previous = line

        self._flush_paragraph()
        self._close_heading_run()
        if self.outline and pending:
            # Outline entries whose text was not found on the page still open the section
            # that applies to what follows.
            self.section = pending[-1].title[:300]


def _assemble_pdf_segments(
    pages: list[list[_Line]], outline: list[_OutlineEntry], body: float
) -> tuple[list[Segment], int, int]:
    """Segments, number of headings seen, and (for outlines) how many entries matched."""
    assembler = _PdfAssembler(outline, body)
    for page_number, lines in enumerate(pages, start=1):
        assembler.add_page(page_number, lines)
    return assembler.segments, assembler.headings, assembler.matched


# --- DOCX ------------------------------------------------------------------------


def _parse_docx(path: Path) -> ParsedDocument:
    import docx
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        document = docx.Document(str(path))
    except Exception as exc:
        raise ParseError(
            "The file could not be read as a DOCX document (it may be corrupt)."
        ) from exc

    segments: list[Segment] = []
    section = ""
    title = ""
    headings = 0

    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            paragraph = Paragraph(child, document)
            text = clean_text(paragraph.text)
            if not text:
                continue
            try:
                style = (paragraph.style.name or "").strip().lower()
            except Exception:
                style = ""
            if style == "title" or re.match(r"^heading \d", style):
                section = text[:300]
                title = title or section
                headings += 1
            else:
                segments.append(Segment(1, section, text))
        elif child.tag == qn("w:tbl"):
            table = Table(child, document)
            for row in table.rows:
                cells: list[str] = []
                for cell in row.cells:
                    cell_text = clean_text(cell.text).replace("\n", " ")
                    if cell_text and (not cells or cells[-1] != cell_text):  # merged cells repeat
                        cells.append(cell_text)
                if cells:
                    segments.append(Segment(1, section, " | ".join(cells)))

    if not segments:
        raise EmptyDocumentError("The document contains no extractable text.")

    core = document.core_properties
    created = core.created.date() if core.created else None
    properties = DocumentProperties(title=clean_text(core.title or title)[:300], created=created)
    return ParsedDocument(segments, properties, 1, {"headings": headings})


# --- Text decoding shared by Markdown, HTML and TXT ---------------------------------


def decode_text(raw: bytes) -> str:
    """Decode a text file: BOM-aware UTF-8/16/32, then strict UTF-8, then Windows-1252."""
    if raw.startswith(b"\xef\xbb\xbf"):
        return raw[3:].decode("utf-8", errors="replace")
    if raw.startswith((b"\xff\xfe\x00\x00", b"\x00\x00\xfe\xff")):
        return raw.decode("utf-32", errors="replace")
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        return raw.decode("utf-16", errors="replace")
    if b"\x00" in raw[:8192]:
        raise ParseError("The file does not look like a text document.")
    try:
        return raw.decode("utf-8")
    except UnicodeDecodeError:
        return raw.decode("cp1252", errors="replace")


# --- Markdown --------------------------------------------------------------------

_ATX_HEADING = re.compile(r"^ {0,3}(#{1,6})[ \t]+(.+?)[ \t]*#*[ \t]*$")
_SETEXT_UNDERLINE = re.compile(r"^ {0,3}(=+|-+)[ \t]*$")
_FENCE = re.compile(r"^ {0,3}(`{3,}|~{3,})")
_HORIZONTAL_RULE = re.compile(r"^ {0,3}([-*_])(?: ?\1){2,}[ \t]*$")
_MD_IMAGE = re.compile(r"!\[([^\]]*)\]\([^)]*\)")
_MD_LINK = re.compile(r"\[([^\]]+)\]\(([^)\s]+)[^)]*\)")
_HTML_COMMENT = re.compile(r"<!--.*?-->", re.DOTALL)
_FRONT_MATTER_KEY = re.compile(r"^([A-Za-z_][\w -]*):\s*(.*)$")


def _split_front_matter(lines: list[str]) -> tuple[dict[str, str], list[str]]:
    if not lines or lines[0].strip() != "---":
        return {}, lines
    for end in range(1, min(len(lines), 60)):
        if lines[end].strip() in {"---", "..."}:
            props: dict[str, str] = {}
            for raw in lines[1:end]:
                match = _FRONT_MATTER_KEY.match(raw.strip())
                if match:
                    props[match.group(1).strip().lower().replace(" ", "_")] = (
                        match.group(2).strip().strip("\"'")
                    )
            return props, lines[end + 1 :]
    return {}, lines


def _iso_date(value: str) -> date | None:
    match = re.match(r"^(\d{4})-(\d{2})-(\d{2})", value.strip())
    if not match:
        return None
    try:
        return date(int(match.group(1)), int(match.group(2)), int(match.group(3)))
    except ValueError:
        return None


def _markdown_inline(text: str) -> str:
    text = _HTML_COMMENT.sub("", text)
    text = _MD_IMAGE.sub(r"\1", text)
    return _MD_LINK.sub(r"\1 (\2)", text)


def _parse_markdown(raw: bytes) -> ParsedDocument:
    front, lines = _split_front_matter(
        decode_text(raw).replace("\r\n", "\n").replace("\r", "\n").split("\n")
    )
    segments: list[Segment] = []
    section = ""
    title = ""
    block: list[str] = []
    headings = 0

    def flush() -> None:
        if not block:
            return
        structured = any(
            _BULLET.match(ln.strip()) or ln.lstrip().startswith(("|", ">")) for ln in block
        )
        text = "\n".join(ln.strip() for ln in block) if structured else _join_wrapped(block)
        text = clean_text(_markdown_inline(text))
        if text:
            segments.append(Segment(1, section, text))
        block.clear()

    index = 0
    while index < len(lines):
        line = lines[index]
        fence = _FENCE.match(line)
        if fence:
            flush()
            marker = fence.group(1)[0]
            code: list[str] = []
            index += 1
            while index < len(lines) and not lines[index].lstrip().startswith(marker * 3):
                code.append(lines[index])
                index += 1
            text = clean_text("\n".join(code))
            if text:
                segments.append(Segment(1, section, text))
            index += 1
            continue

        atx = _ATX_HEADING.match(line)
        following = lines[index + 1] if index + 1 < len(lines) else ""
        if atx:
            flush()
            section = clean_text(_markdown_inline(atx.group(2)))[:300]
            title = title or section
            headings += 1
        elif (
            line.strip()
            and not _BULLET.match(line.strip())
            and _SETEXT_UNDERLINE.match(following)
            and not block
        ):
            flush()
            section = clean_text(_markdown_inline(line))[:300]
            title = title or section
            headings += 1
            index += 1  # skip the underline
        elif not line.strip() or _HORIZONTAL_RULE.match(line):
            flush()
        else:
            block.append(line)
        index += 1
    flush()

    if not segments:
        raise EmptyDocumentError("The document contains no extractable text.")
    properties = DocumentProperties(
        title=clean_text(front.get("title", "") or title)[:300],
        created=_iso_date(front.get("date", "") or front.get("created", "")),
        department=front.get("department", ""),
        document_type=front.get("document_type", "") or front.get("type", ""),
    )
    return ParsedDocument(segments, properties, 1, {"headings": headings})


# --- HTML ------------------------------------------------------------------------

_HTML_DROP = {
    "script",
    "style",
    "noscript",
    "template",
    "svg",
    "canvas",
    "iframe",
    "button",
    "select",
    "input",
    "head",
}
_HTML_HEADINGS = {"h1", "h2", "h3", "h4", "h5", "h6"}
_HTML_INLINE = {
    "a", "abbr", "acronym", "b", "bdi", "bdo", "big", "cite", "code", "data", "del", "dfn", "em", "font",
    "i", "ins", "kbd", "label", "mark", "q", "s", "samp", "small", "span", "strike", "strong", "sub",
    "sup", "time", "tt", "u", "var", "wbr",
}  # fmt: skip


def _collapse(text: str) -> str:
    """Collapse every run of whitespace, newlines included, to one space."""
    return clean_text(re.sub(r"\s+", " ", text))


_HTML_DATE_META = (
    "article:published_time",
    "date",
    "dcterms.created",
    "dcterms.date",
    "created",
    "pubdate",
)


def _parse_html(raw: bytes) -> ParsedDocument:
    from bs4 import BeautifulSoup, NavigableString, Tag
    from bs4.element import PreformattedString

    if b"\x00" in raw[:8192] and not raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        raise ParseError("The file does not look like an HTML document.")
    soup = BeautifulSoup(raw, "html.parser")

    # The title and date live in <head>, which is dropped below, so read them first.
    title = _collapse(soup.title.get_text(" ")) if soup.title else ""
    created: date | None = None
    for meta in soup.find_all("meta"):
        name = (meta.get("name") or meta.get("property") or "").lower()
        if name in _HTML_DATE_META and meta.get("content"):
            created = created or _iso_date(str(meta["content"]))

    for tag in soup.find_all(_HTML_DROP):
        tag.decompose()

    root = soup.find("main") or soup.find("article") or soup.body or soup
    segments: list[Segment] = []
    section = ""
    inline: list[str] = []
    headings = 0

    def flush() -> None:
        text = _collapse(" ".join(inline))
        inline.clear()
        if text:
            segments.append(Segment(1, section, text))

    def walk(node: Tag) -> None:
        nonlocal section, headings
        for child in node.children:
            if isinstance(child, NavigableString):
                if not isinstance(child, PreformattedString):  # comments, doctype, CDATA, ...
                    inline.append(str(child))
                continue
            if not isinstance(child, Tag):
                continue
            name = (child.name or "").lower()
            if name in _HTML_HEADINGS:
                flush()
                heading = _collapse(child.get_text(" "))
                if heading:
                    section = heading[:300]
                    headings += 1
            elif name == "br":
                inline.append(" ")
            elif name == "pre":
                flush()
                text = clean_text(child.get_text())
                if text:
                    segments.append(Segment(1, section, text))
            elif name == "tr":
                flush()
                cells = [_collapse(c.get_text(" ")) for c in child.find_all(["td", "th"])]
                row = " | ".join(c for c in cells if c)
                if row:
                    segments.append(Segment(1, section, row))
            elif name in _HTML_INLINE:
                walk(child)
            elif name in {"nav", "footer", "aside"} and root is soup.body:
                continue  # page chrome on documents without a <main>/<article>
            else:
                flush()  # any other element is a block: its text starts a new paragraph
                walk(child)
                flush()

    walk(root)
    flush()

    if not segments:
        raise EmptyDocumentError("The document contains no extractable text.")
    return ParsedDocument(
        segments,
        DocumentProperties(title=title[:300] or section, created=created),
        1,
        {"headings": headings},
    )


# --- Plain text --------------------------------------------------------------------

_TXT_UNDERLINE = re.compile(r"^(={3,}|-{3,}|~{3,})\s*$")


def _txt_is_heading(line: str) -> bool:
    text = line.strip()
    return (
        3 <= len(text) <= 80
        and _letters(text) >= 3
        and text == text.upper()
        and not text.endswith((".", ",", ";", ":"))
        and not _BULLET.match(text)
    )


def _parse_text(raw: bytes) -> ParsedDocument:
    lines = decode_text(raw).replace("\r\n", "\n").replace("\r", "\n").split("\n")
    segments: list[Segment] = []
    section = ""
    title = ""
    block: list[str] = []
    headings = 0

    def flush() -> None:
        if not block:
            return
        structured = any(ln[:1] in {" ", "\t"} or _BULLET.match(ln.strip()) for ln in block)
        text = "\n".join(ln.rstrip() for ln in block) if structured else _join_wrapped(block)
        text = clean_text(text)
        if text:
            segments.append(Segment(1, section, text))
        block.clear()

    for index, line in enumerate(lines):
        previous_blank = index == 0 or not lines[index - 1].strip()
        if not line.strip():
            flush()
        elif _TXT_UNDERLINE.match(line) and block and len(block) == 1:
            section = clean_text(block.pop())[:300]  # setext-style heading
            title = title or section
            headings += 1
        elif _TXT_UNDERLINE.match(line):
            continue
        elif previous_blank and not block and _txt_is_heading(line):
            section = clean_text(line)[:300]
            title = title or section
            headings += 1
        else:
            block.append(line)
    flush()

    if not segments:
        raise EmptyDocumentError("The document contains no extractable text.")
    return ParsedDocument(
        segments, DocumentProperties(title=title[:300]), 1, {"headings": headings}
    )


# --- Dispatch ----------------------------------------------------------------------


def _check_docx_container(path: Path) -> None:
    """Reject DOCX files whose zip structure looks like a decompression bomb."""
    try:
        with zipfile.ZipFile(path) as archive:
            infos = archive.infolist()
            if len(infos) > 20_000 or sum(i.file_size for i in infos) > 512 * 1024 * 1024:
                raise ParseError("The DOCX file is too large when unpacked.")
    except zipfile.BadZipFile as exc:
        raise ParseError(
            "The file could not be read as a DOCX document (it may be corrupt)."
        ) from exc


def parse_document(path: Path, extension: str) -> ParsedDocument:
    """Parse the file at ``path``. ``extension`` (e.g. ``.pdf``) selects the parser."""
    ext = extension.lower()
    if ext == ".pdf":
        return _parse_pdf(path)
    if ext == ".docx":
        _check_docx_container(path)
        return _parse_docx(path)
    if ext in {".md", ".markdown"}:
        return _parse_markdown(path.read_bytes())
    if ext in {".html", ".htm"}:
        return _parse_html(path.read_bytes())
    if ext == ".txt":
        return _parse_text(path.read_bytes())
    raise UnsupportedFormatError(f"Unsupported file type '{ext[:12]}'.")
