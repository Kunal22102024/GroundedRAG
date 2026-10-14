"""Parsers: page numbers, headings, cleanup, failure messages and every supported format."""

from __future__ import annotations

import contextlib
import time
from datetime import date, datetime
from pathlib import Path

import pytest

from app.core import parsers
from app.core.parsers import (
    EmptyDocumentError,
    ParseError,
    ScannedPdfError,
    UnsupportedFormatError,
    parse_document,
)
from tests.pdf_factory import ZERO_PAGE_PDF, make_pdf

LOREM = (
    "The quick brown fox jumps over the lazy dog while the committee reviews the quarterly "
    "retention policy and records its findings in the central register."
)
_WORDS = [
    "archive",
    "balance",
    "carrier",
    "ledger",
    "harbor",
    "invoice",
    "journal",
    "kettle",
    "lantern",
    "meadow",
    "network",
    "orchard",
    "pantry",
    "quarry",
    "ribbon",
    "saddle",
    "timber",
    "upland",
    "vessel",
    "willow",
    "anchor",
    "bridge",
    "cobalt",
    "dynamo",
    "ember",
    "falcon",
    "garnet",
    "hollow",
    "island",
    "jasper",
    "keystone",
    "lagoon",
    "marble",
    "nectar",
    "opal",
    "prairie",
    "quartz",
    "raven",
    "summit",
    "tundra",
    "umber",
    "velvet",
    "walnut",
    "yonder",
    "zephyr",
    "bramble",
    "cinder",
    "dusk",
    "echo",
    "fjord",
    "glacier",
    "harvest",
    "ivory",
    "juniper",
    "kindle",
    "lumen",
    "mosaic",
    "nimbus",
    "onyx",
    "pebble",
    "quiver",
    "russet",
    "sable",
    "thistle",
]


def body(n: int, words: int = 24) -> str:
    """Page text that differs from every other page (so nothing looks like repeated page chrome)."""
    return " ".join(_WORDS[(n * 11 + i * 3) % len(_WORDS)] for i in range(words)) + "."


def write(tmp_path: Path, name: str, content: bytes | str) -> Path:
    path = tmp_path / name
    path.write_bytes(content if isinstance(content, bytes) else content.encode("utf-8"))
    return path


def pdf_file(tmp_path: Path, pages: list, **kwargs: object) -> Path:
    return write(tmp_path, "doc.pdf", make_pdf(pages, **kwargs))  # type: ignore[arg-type]


# --- PDF: pages ------------------------------------------------------------------------


def test_pdf_segments_carry_the_real_physical_page_number(tmp_path: Path) -> None:
    pages = [[("p", f"Marker{n} {body(n)}")] for n in range(1, 6)]
    doc = parse_document(pdf_file(tmp_path, pages), ".pdf")
    assert doc.page_count == 5
    for segment in doc.segments:
        marker = int(segment.text.split()[0].removeprefix("Marker"))
        assert segment.page == marker, f"text from page {marker} was labelled page {segment.page}"
    assert {s.page for s in doc.segments} == {1, 2, 3, 4, 5}


def test_pdf_blank_pages_in_the_middle_do_not_shift_page_numbers(tmp_path: Path) -> None:
    pages = [[("p", f"First {body(1)}")], None, None, [("p", f"Fourth {body(4)}")]]
    doc = parse_document(pdf_file(tmp_path, pages), ".pdf")
    by_marker = {s.text.split()[0]: s.page for s in doc.segments}
    assert by_marker == {"First": 1, "Fourth": 4}
    assert doc.stats["pages_without_text"] == 2


def test_pdf_paragraphs_are_separated_by_visible_gaps(tmp_path: Path) -> None:
    pages = [[("p", f"Alpha {body(1)}"), ("p", f"Bravo {body(2)}"), ("p", f"Charlie {body(3)}")]]
    doc = parse_document(pdf_file(tmp_path, pages), ".pdf")
    starts = [s.text.split()[0] for s in doc.segments]
    assert starts == ["Alpha", "Bravo", "Charlie"]


def test_pdf_hyphenated_line_breaks_are_repaired(tmp_path: Path) -> None:
    text = "The administra-\ntive procedure for reimburse-\nment is described below in detail."
    doc = parse_document(pdf_file(tmp_path, [[("p", text)]]), ".pdf")
    joined = " ".join(s.text for s in doc.segments)
    assert "administrative" in joined and "reimbursement" in joined


# --- PDF: headings and sections ----------------------------------------------------------


def test_pdf_headings_found_by_font_size_become_sections(tmp_path: Path) -> None:
    pages = [
        [("h", "Purpose"), ("p", f"Aim {body(1)}"), ("h", "Scope"), ("p", f"Reach {body(2)}")],
        [("p", f"Continues {body(3)}"), ("h", "Safeguards"), ("p", f"Rules {body(4)}")],
    ]
    doc = parse_document(pdf_file(tmp_path, pages), ".pdf")
    sections = {s.text.split()[0]: s.section for s in doc.segments}
    assert sections == {
        "Aim": "Purpose",
        "Reach": "Scope",
        "Continues": "Scope",
        "Rules": "Safeguards",
    }
    assert [s.page for s in doc.segments if s.text.startswith("Continues")] == [2]


def test_pdf_bold_short_lines_are_headings(tmp_path: Path) -> None:
    pages = [[("b", "Retention Period"), ("p", f"Body {body(1)}")]]
    doc = parse_document(pdf_file(tmp_path, pages), ".pdf")
    assert doc.segments[0].section == "Retention Period"


def test_pdf_outline_provides_sections_when_it_has_enough_entries(tmp_path: Path) -> None:
    pages = [
        [("b", name), ("p", f"{name}Body {body(i)}")]
        for i, name in enumerate(("Overview", "Safety", "Charging", "Service"), 1)
    ]
    outline = [(1, "Overview", 1), (1, "Safety", 2), (2, "Charging", 3), (1, "Service", 4)]
    doc = parse_document(pdf_file(tmp_path, pages, outline=outline), ".pdf")
    assert doc.stats["outline_entries"] == 4
    assert doc.stats["outline_entries_matched"] == 4
    sections = {s.text.split()[0]: (s.section, s.page) for s in doc.segments}
    assert sections["OverviewBody"] == ("Overview", 1)
    assert sections["ChargingBody"] == ("Charging", 3)
    assert sections["ServiceBody"] == ("Service", 4)


def test_pdf_with_too_few_outline_entries_falls_back_to_font_headings(tmp_path: Path) -> None:
    pages = [[("h", "Real Heading"), ("p", f"Body {body(1)}")]]
    doc = parse_document(pdf_file(tmp_path, pages, outline=[(1, "Only One", 1)]), ".pdf")
    assert doc.stats["outline_entries"] == 0
    assert doc.segments[0].section == "Real Heading"


def test_pdf_running_headers_and_footers_are_removed(tmp_path: Path) -> None:
    pages = [[("p", f"Unique{n} {body(n, 60)}")] for n in range(1, 7)]
    doc = parse_document(
        pdf_file(tmp_path, pages, header="ACME CONFIDENTIAL DRAFT", footer=True), ".pdf"
    )
    text = " ".join(s.text for s in doc.segments)
    assert "ACME CONFIDENTIAL" not in text
    assert "Page 3" not in text
    assert doc.stats["header_footer_lines_removed"] >= 6
    assert all(f"Unique{n}" in text for n in range(1, 7)), "body text must survive"


def test_a_document_of_near_identical_pages_is_not_mistaken_for_page_chrome(tmp_path: Path) -> None:
    """Forms and templates repeat. Removing 'repeated lines' must never empty the document."""
    pages = [[("p", f"Form{n} {LOREM}")] for n in range(1, 7)]
    doc = parse_document(pdf_file(tmp_path, pages), ".pdf")
    assert doc.stats["header_footer_lines_removed"] == 0
    assert len(doc.segments) == 6
    assert all("committee reviews the quarterly" in s.text for s in doc.segments)


def test_pdf_header_removal_needs_enough_pages_to_be_sure(tmp_path: Path) -> None:
    pages = [[("p", f"Unique{n} {body(n)}")] for n in range(1, 3)]
    doc = parse_document(pdf_file(tmp_path, pages, header="SHORT DOC HEADER"), ".pdf")
    assert "SHORT DOC HEADER" in " ".join(s.text for s in doc.segments)


# --- PDF: columns, tables and page labels -----------------------------------------------------------


LEFT_COLUMN = " ".join(f"leftword{n}" for n in range(1, 60)) + " LEFTEND."
RIGHT_COLUMN = " ".join(f"rightword{n}" for n in range(1, 60)) + " RIGHTEND."


def test_two_column_pages_are_read_column_by_column(tmp_path: Path) -> None:
    """Side-by-side columns must not be stitched into lines that mix two sentences."""
    doc = parse_document(pdf_file(tmp_path, [[("cols", f"{LEFT_COLUMN}||{RIGHT_COLUMN}")]]), ".pdf")
    assert doc.stats["two_column_pages"] == 1
    for segment in doc.segments:
        assert not ("leftword" in segment.text and "rightword" in segment.text), segment.text[:120]
    joined = " ".join(s.text for s in doc.segments)
    assert joined.index("LEFTEND") < joined.index("rightword1 "), (
        "the whole left column comes first"
    )
    assert "leftword30 leftword31" in joined and "rightword30 rightword31" in joined, (
        "words stay in order"
    )


def test_a_heading_across_both_columns_labels_everything_below_it(tmp_path: Path) -> None:
    pages = [
        [("h", "Shared Heading Over Both Columns"), ("cols", f"{LEFT_COLUMN}||{RIGHT_COLUMN}")]
    ]
    doc = parse_document(pdf_file(tmp_path, pages), ".pdf")
    assert doc.stats["two_column_pages"] == 1
    assert {s.section for s in doc.segments} == {"Shared Heading Over Both Columns"}


def test_a_table_of_short_cells_is_not_split_into_columns(tmp_path: Path) -> None:
    rows = [("row", f"E{n}||Disk full {n}||Check the cable") for n in range(1, 9)]
    doc = parse_document(pdf_file(tmp_path, [rows]), ".pdf")
    assert doc.stats["two_column_pages"] == 0
    text = " ".join(s.text for s in doc.segments)
    for n in range(1, 9):
        assert f"E{n} Disk full {n} Check the cable" in text, "each row stays together, in order"


def test_single_column_pages_are_unaffected(tmp_path: Path) -> None:
    doc = parse_document(pdf_file(tmp_path, [[("p", body(1, 80)), ("p", body(2, 80))]]), ".pdf")
    assert doc.stats["two_column_pages"] == 0 and len(doc.segments) == 2


def test_page_labels_are_removed_even_in_short_documents(tmp_path: Path) -> None:
    pages = [
        [("p", f"First page text. {body(1, 40)}")],
        [("p", f"Second page text. {body(2, 40)}")],
    ]
    doc = parse_document(pdf_file(tmp_path, pages, footer=True), ".pdf")
    text = " ".join(s.text for s in doc.segments)
    assert "Page 1" not in text and "Page 2" not in text
    assert "First page text." in text and "Second page text." in text


@pytest.mark.parametrize(
    ("label", "is_label"),
    [
        ("Page 3", True), ("page 12 of 40", True), ("Pg. 7", True), ("p. 9", True), ("3 of 12", True),
        ("- 4 -", True), ("42", True), ("\u2014 17 \u2014", True),
        ("Page limits apply to every submission", False), ("Section 3", False), ("Table 4", False),
        ("12 apples", False), ("Chapter 2 of the contract", False),
    ],
)  # fmt: skip
def test_page_label_pattern(label: str, is_label: bool) -> None:
    assert bool(parsers._PAGE_LABEL.match(label)) is is_label


# --- PDF: properties --------------------------------------------------------------------


def test_pdf_properties_title_and_creation_date(tmp_path: Path) -> None:
    doc = parse_document(
        pdf_file(
            tmp_path,
            [[("p", body(1))]],
            title="Email Management Policy",
            created=datetime(2025, 4, 7, 21, 56),
        ),
        ".pdf",
    )
    assert doc.properties.title == "Email Management Policy"
    assert doc.properties.created == date(2025, 4, 7)


def test_pdf_date_parsing() -> None:
    assert parsers._parse_pdf_date("D:20260201143800Z00'00'") == date(2026, 2, 1)
    assert parsers._parse_pdf_date("D:2012") == date(2012, 1, 1)
    assert parsers._parse_pdf_date("garbage") is None
    assert parsers._parse_pdf_date(None) is None
    assert parsers._parse_pdf_date("D:20261399") is None


# --- PDF: documents that cannot be read ---------------------------------------------------


def test_scanned_pdf_fails_with_the_ocr_message(tmp_path: Path) -> None:
    with pytest.raises(ScannedPdfError) as caught:
        parse_document(pdf_file(tmp_path, ["image", "image", "image"]), ".pdf")
    assert "scanned PDF, OCR not supported" in str(caught.value)


def test_pdf_with_only_blank_pages_is_treated_as_scanned(tmp_path: Path) -> None:
    with pytest.raises(ScannedPdfError, match="OCR not supported"):
        parse_document(pdf_file(tmp_path, [None, None]), ".pdf")


def test_pdf_that_is_mostly_scanned_is_rejected_but_mostly_text_is_kept(tmp_path: Path) -> None:
    mostly_scans = ["image"] * 9 + [[("p", body(1))]]
    with pytest.raises(ScannedPdfError):
        parse_document(pdf_file(tmp_path, mostly_scans), ".pdf")
    mostly_text = [[("p", f"Page{n} {body(n)}")] for n in range(1, 5)] + [None]
    assert parse_document(pdf_file(tmp_path, mostly_text), ".pdf").stats["pages_without_text"] == 1


def test_pdf_without_pages_is_an_empty_document(tmp_path: Path) -> None:
    with pytest.raises(EmptyDocumentError, match="no pages"):
        parse_document(write(tmp_path, "zero.pdf", ZERO_PAGE_PDF), ".pdf")


def test_corrupt_pdf_gives_a_clear_message(tmp_path: Path) -> None:
    with pytest.raises(ParseError, match="could not be read as a PDF"):
        parse_document(
            write(tmp_path, "bad.pdf", b"%PDF-1.4 this is not really a pdf at all" * 5), ".pdf"
        )


def test_parse_errors_never_contain_document_text(tmp_path: Path) -> None:
    secret = "TOPSECRETPHRASE"
    path = write(tmp_path, "bad.pdf", f"%PDF-1.4 {secret}".encode() * 4)
    with pytest.raises(ParseError) as caught:
        parse_document(path, ".pdf")
    assert secret not in str(caught.value)


# --- DOCX -------------------------------------------------------------------------------


def make_docx(tmp_path: Path, build) -> Path:  # noqa: ANN001
    import docx

    document = docx.Document()
    build(document)
    path = tmp_path / "doc.docx"
    document.save(path)
    return path


def test_docx_headings_paragraphs_and_tables(tmp_path: Path) -> None:
    def build(document) -> None:  # noqa: ANN001
        document.add_heading("Leave Policy", level=1)
        document.add_paragraph("Employees receive 25 days of leave.")
        document.add_heading("Carry-over", level=2)
        document.add_paragraph("Up to five days may be carried over.")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text, table.cell(0, 1).text = "Grade", "Days"
        table.cell(1, 0).text, table.cell(1, 1).text = "Senior", "30"
        document.core_properties.created = datetime(2024, 3, 9, 10, 0)

    doc = parse_document(make_docx(tmp_path, build), ".docx")
    assert [(s.section, s.text) for s in doc.segments] == [
        ("Leave Policy", "Employees receive 25 days of leave."),
        ("Carry-over", "Up to five days may be carried over."),
        ("Carry-over", "Grade | Days"),
        ("Carry-over", "Senior | 30"),
    ]
    assert all(s.page == 1 for s in doc.segments)
    assert doc.properties.created == date(2024, 3, 9)


def test_docx_with_no_text_is_an_empty_document(tmp_path: Path) -> None:
    with pytest.raises(EmptyDocumentError, match="no extractable text"):
        parse_document(make_docx(tmp_path, lambda d: d.add_paragraph("   ")), ".docx")


def test_corrupt_docx_gives_a_clear_message(tmp_path: Path) -> None:
    with pytest.raises(ParseError, match="DOCX"):
        parse_document(write(tmp_path, "bad.docx", b"PK\x03\x04 definitely not a zip"), ".docx")


# --- Markdown ---------------------------------------------------------------------------

MARKDOWN = """---
title: Expense Policy
department: Finance
type: Policy
date: 2025-06-30
---

# Expenses

Claims must be filed within
30 days of purchase.

## Limits

- Meals: 50 per day
- Travel: economy only

```python
LIMIT = 50
```

Setext Heading
==============

See [the portal](https://example.com/portal) for ![diagram](d.png) details.

---

Final paragraph here.
"""


def test_markdown_structure_front_matter_and_cleanup(tmp_path: Path) -> None:
    doc = parse_document(write(tmp_path, "p.md", MARKDOWN), ".md")
    by_text = {s.text: s.section for s in doc.segments}
    assert by_text["Claims must be filed within 30 days of purchase."] == "Expenses"
    assert by_text["- Meals: 50 per day\n- Travel: economy only"] == "Limits"
    assert by_text["LIMIT = 50"] == "Limits", "code blocks are kept verbatim, without fences"
    assert (
        by_text["See the portal (https://example.com/portal) for diagram details."]
        == "Setext Heading"
    )
    assert by_text["Final paragraph here."] == "Setext Heading"
    assert "---" not in by_text, "a horizontal rule is not text"
    assert doc.properties.title == "Expense Policy"
    assert (doc.properties.department, doc.properties.document_type) == ("Finance", "Policy")
    assert doc.properties.created == date(2025, 6, 30)


def test_markdown_decodes_bom_crlf_and_windows_1252(tmp_path: Path) -> None:
    bom = write(tmp_path, "a.md", b"\xef\xbb\xbf# Title\r\n\r\nCaf\xc3\xa9 text here.\r\n")
    assert parse_document(bom, ".md").segments[0].text == "Café text here."
    cp1252 = write(tmp_path, "b.md", b"# T\n\nR\xe9sum\xe9 text here.\n")
    assert parse_document(cp1252, ".md").segments[0].text == "Résumé text here."


def test_empty_markdown_is_an_empty_document(tmp_path: Path) -> None:
    with pytest.raises(EmptyDocumentError):
        parse_document(write(tmp_path, "e.md", "# Only a heading\n"), ".md")


# --- HTML -------------------------------------------------------------------------------

HTML = """<!doctype html><html><head><title>Network Spec</title>
<meta name="date" content="2024-11-02"><style>p{color:red}</style><script>var secret = 1;</script></head>
<body><nav>Home | About | Contact</nav>
<main><h1>Timeouts</h1><p>The retry limit is <b>3</b> attempts,
with backoff.</p>
<h2>Errors</h2><ul><li>ERR_404_NET: gateway unreachable</li></ul>
<table><tr><th>Code</th><th>Meaning</th></tr><tr><td>E1</td><td>Disk full</td></tr></table>
<pre>retry --max 3
--backoff</pre></main><footer>Copyright</footer></body></html>"""


def test_html_structure_and_noise_removal(tmp_path: Path) -> None:
    doc = parse_document(write(tmp_path, "s.html", HTML), ".html")
    pairs = [(s.section, s.text) for s in doc.segments]
    assert ("Timeouts", "The retry limit is 3 attempts, with backoff.") in pairs
    assert ("Errors", "ERR_404_NET: gateway unreachable") in pairs
    assert ("Errors", "Code | Meaning") in pairs and ("Errors", "E1 | Disk full") in pairs
    assert ("Errors", "retry --max 3\n--backoff") in pairs
    everything = " ".join(t for _, t in pairs)
    for noise in ("secret", "color:red", "Home | About", "Copyright"):
        assert noise not in everything
    assert doc.properties.title == "Network Spec"
    assert doc.properties.created == date(2024, 11, 2)


def test_html_without_main_drops_navigation_chrome(tmp_path: Path) -> None:
    html = "<html><body><nav>Menu stuff</nav><h1>Real</h1><p>Actual content lives here.</p></body></html>"
    doc = parse_document(write(tmp_path, "n.html", html), ".html")
    assert [s.text for s in doc.segments] == ["Actual content lives here."]


def test_html_entities_and_encodings(tmp_path: Path) -> None:
    doc = parse_document(
        write(tmp_path, "e.html", "<p>Fish &amp; chips &mdash; caf\xe9</p>".encode("latin-1")),
        ".html",
    )
    assert "Fish & chips" in doc.segments[0].text


def test_html_that_is_binary_or_empty_is_rejected(tmp_path: Path) -> None:
    with pytest.raises(ParseError):
        parse_document(write(tmp_path, "b.html", b"\x00\x01\x02" * 50), ".html")
    with pytest.raises(EmptyDocumentError):
        parse_document(
            write(tmp_path, "e.html", "<html><body><script>x()</script></body></html>"), ".html"
        )


# --- Plain text ---------------------------------------------------------------------------


def test_text_paragraphs_headings_and_wrapped_lines(tmp_path: Path) -> None:
    text = (
        "PURPOSE\n\nThis policy explains how long\nrecords are kept before disposal.\n\n"
        "RETENTION PERIODS\n\n  Email: 3 years\n  Contracts: 7 years\n\n"
        "Setext Title\n============\nBody after underline here.\n"
    )
    doc = parse_document(write(tmp_path, "p.txt", text), ".txt")
    pairs = [(s.section, s.text) for s in doc.segments]
    assert ("PURPOSE", "This policy explains how long records are kept before disposal.") in pairs
    assert ("RETENTION PERIODS", "Email: 3 years\nContracts: 7 years") in pairs, (
        "list lines keep their line breaks"
    )
    assert ("Setext Title", "Body after underline here.") in pairs


def test_text_decodes_utf16_and_rejects_binary(tmp_path: Path) -> None:
    utf16 = write(tmp_path, "u.txt", "Hello wide world, this is text.".encode("utf-16"))
    assert parse_document(utf16, ".txt").segments[0].text == "Hello wide world, this is text."
    with pytest.raises(ParseError, match="text document"):
        parse_document(write(tmp_path, "b.txt", b"abc\x00\x01\x02def" * 30), ".txt")


def test_whitespace_only_text_is_an_empty_document(tmp_path: Path) -> None:
    with pytest.raises(EmptyDocumentError):
        parse_document(write(tmp_path, "w.txt", "  \n\n \t \n"), ".txt")


# --- Dispatch and text cleaning -------------------------------------------------------------


def test_unsupported_extension(tmp_path: Path) -> None:
    with pytest.raises(UnsupportedFormatError, match="Unsupported file type"):
        parse_document(write(tmp_path, "x.exe", b"MZ"), ".exe")


def test_extension_matching_is_case_insensitive(tmp_path: Path) -> None:
    assert parse_document(write(tmp_path, "a.TXT", "Some plain text content."), ".TXT").segments


def test_clean_text_normalizes_ligatures_invisible_and_control_characters() -> None:
    dirty = "oﬃce­  work​flow\x07  tab\there"
    assert parsers.clean_text(dirty) == "office workflow tab here"


# --- Hostile input: parsing time grows with the length, not with its square ----------------------------------


@pytest.mark.parametrize(
    "text",
    [
        "[" * 200_000,
        "![" * 100_000,
        "[a](" * 50_000,
        "](" * 200_000,
        "<!--" * 100_000,
        "# a" + " " * 100_000 + "b",
        "# " + "# " * 50_000,
    ],
    ids=["brackets", "images", "links", "link_openers", "comments", "padded_heading", "hashes"],
)
def test_markdown_parses_hostile_input_in_linear_time(tmp_path: Path, text: str) -> None:
    """Each of these took 8 s to 50 s or more per 200 KB with the old regular expressions."""
    path = write(tmp_path, "h.md", text)
    start = time.perf_counter()
    with contextlib.suppress(EmptyDocumentError):  # e.g. a document that is only a heading
        parse_document(path, ".md")
    assert time.perf_counter() - start < 3.0


def test_text_with_hundreds_of_thousands_of_short_lines_parses_quickly(tmp_path: Path) -> None:
    path = write(tmp_path, "l.txt", "\n".join(["word"] * 400_000))
    start = time.perf_counter()
    parse_document(path, ".txt")
    assert time.perf_counter() - start < 3.0


def test_wrapped_lines_are_joined_and_hyphenation_is_repaired() -> None:
    join = parsers._join_wrapped
    assert join(["alpha", "  beta ", "", "gamma"]) == "alpha beta gamma"
    assert join(["an exam-", "ple of it"]) == "an example of it"
    assert join(["keep a dash -", "after a space"]) == "keep a dash - after a space"
    assert join(["Upper-", "Case stays"]) == "Upper- Case stays"
    assert join([]) == ""


def test_markdown_links_images_and_comments_are_flattened(tmp_path: Path) -> None:
    md = (
        'Read [the guide](https://x.io/g "Title") and ![logo](logo.png) now.\n\n'
        "[a](u)![b](v)[c](w)\n\n"
        "A [plain] (bracket) stays, and so does [unfinished](\n\n"
        "Before <!-- hidden --> after <!-- never closed\n"
    )
    doc = parse_document(write(tmp_path, "l.md", md), ".md")
    assert [s.text for s in doc.segments] == [
        "Read the guide (https://x.io/g) and logo now.",
        "a (u)bc (w)",
        "A [plain] (bracket) stays, and so does [unfinished](",
        "Before after <!-- never closed",
    ]


def test_an_image_with_an_embedded_data_uri_is_dropped_whatever_its_size(tmp_path: Path) -> None:
    md = f"Chart: ![sales](data:image/png;base64,{'A' * 300_000}) ends here.\n"
    text = parse_document(write(tmp_path, "d.md", md), ".md").segments[0].text
    assert text == "Chart: sales ends here."


def test_a_lone_identifier_line_in_a_text_file_is_content_not_a_heading(tmp_path: Path) -> None:
    text = "OVERVIEW\n\nERR_404_NET\n\nSECTION 3\n\nBody under section three is here.\n"
    doc = parse_document(write(tmp_path, "i.txt", text), ".txt")
    pairs = [(s.section, s.text) for s in doc.segments]
    assert ("OVERVIEW", "ERR_404_NET") in pairs
    assert ("SECTION 3", "Body under section three is here.") in pairs


# --- DOCX content controls ---------------------------------------------------------------------------------


def _docx_with_content_controls(tmp_path: Path) -> Path:
    import docx
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    document = docx.Document()
    document.add_heading("Leave Policy", level=1)
    document.add_paragraph("Ordinary paragraph before the controls.")
    body = document.element.body
    sectpr = body[-1]  # the section properties stay last

    block = parse_xml(
        f"<w:sdt {nsdecls('w')}><w:sdtPr/><w:sdtContent>"
        "<w:p><w:r><w:t>Paragraph inside a content control.</w:t></w:r></w:p>"
        "<w:tbl><w:tr><w:tc><w:p><w:r><w:t>Grade</w:t></w:r></w:p></w:tc>"
        "<w:tc><w:p><w:r><w:t>Band</w:t></w:r></w:p></w:tc></w:tr></w:tbl>"
        "</w:sdtContent></w:sdt>"
    )
    nested = parse_xml(
        f"<w:sdt {nsdecls('w')}><w:sdtContent><w:sdt><w:sdtContent>"
        "<w:p><w:r><w:t>Nested control text.</w:t></w:r></w:p>"
        "</w:sdtContent></w:sdt></w:sdtContent></w:sdt>"
    )
    inline = parse_xml(
        f"<w:p {nsdecls('w')}><w:r><w:t xml:space='preserve'>Employee: </w:t></w:r>"
        "<w:sdt><w:sdtContent><w:r><w:t>Ada Lovelace</w:t></w:r></w:sdtContent></w:sdt>"
        "<w:r><w:t xml:space='preserve'> is entitled to 25 days.</w:t></w:r></w:p>"
    )
    for element in (block, nested, inline):
        sectpr.addprevious(element)
    path = tmp_path / "controls.docx"
    document.save(str(path))
    return path


def test_docx_text_inside_content_controls_is_not_dropped(tmp_path: Path) -> None:
    doc = parse_document(_docx_with_content_controls(tmp_path), ".docx")
    pairs = [(s.section, s.text) for s in doc.segments]
    assert ("Leave Policy", "Ordinary paragraph before the controls.") in pairs
    assert ("Leave Policy", "Paragraph inside a content control.") in pairs
    assert ("Leave Policy", "Grade | Band") in pairs, "a table inside a content control"
    assert ("Leave Policy", "Nested control text.") in pairs
    assert ("Leave Policy", "Employee: Ada Lovelace is entitled to 25 days.") in pairs
