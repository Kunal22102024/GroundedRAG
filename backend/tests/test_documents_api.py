"""Documents API: upload validation, role-based visibility, file serving and deletion."""

from __future__ import annotations

import io
import re
from collections.abc import Iterator
from pathlib import Path
from typing import Any

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.config import get_settings
from app.core import bm25, vectorstore
from app.db.models import Chunk, Document
from app.main import app
from app.services import ingestion
from tests.conftest import FakeEmbeddings
from tests.pdf_factory import make_pdf
from tests.test_parsers import body

PASSWORD = "correct horse battery"
MD = b"# Retry Policy\n\nThe retry limit for API timeouts is three attempts with exponential backoff.\n"
Headers = dict[str, str]


@pytest.fixture
def client(
    db: Session,
    upload_dir: Path,
    vectors: Any,
    fake_embeddings: FakeEmbeddings,
    inline_ingestion: list,
) -> Iterator[TestClient]:
    with TestClient(app) as test_client:
        yield test_client


def auth(client: TestClient, email: str) -> Headers:
    response = client.post("/api/auth/signup", json={"email": email, "password": PASSWORD})
    assert response.status_code == 201, response.text
    return {"Authorization": f"Bearer {response.json()['access_token']}"}


@pytest.fixture
def admin(client: TestClient) -> Headers:
    return auth(client, "boss@example.com")  # the first account is the administrator


@pytest.fixture
def user(client: TestClient, admin: Headers) -> Headers:
    return auth(client, "worker@example.com")


@pytest.fixture
def other(client: TestClient, user: Headers) -> Headers:
    return auth(client, "other@example.com")


def upload(
    client: TestClient, headers: Headers, *files: tuple[str, bytes, str], **data: str
) -> Any:
    return client.post(
        "/api/documents",
        headers=headers,
        data=data,
        files=[("files", (name, content, ctype)) for name, content, ctype in files],
    )


def one(
    client: TestClient,
    headers: Headers,
    name: str = "policy.md",
    content: bytes = MD,
    ctype: str = "text/markdown",
    **data: str,
) -> dict[str, Any]:
    response = upload(client, headers, (name, content, ctype), **data)
    assert response.status_code == 202, response.text
    return response.json()["documents"][0]  # type: ignore[no-any-return]


def docx_bytes() -> bytes:
    import docx

    document = docx.Document()
    document.add_heading("Leave Policy", level=1)
    document.add_paragraph("Employees receive twenty five days of paid leave each year.")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# --- Upload ---------------------------------------------------------------------------------------------


def test_uploading_several_files_queues_them_and_they_become_indexed(
    client: TestClient, user: Headers
) -> None:
    response = upload(
        client,
        user,
        ("policy.md", MD, "text/markdown"),
        ("notes.txt", b"Plain notes about the retry limit.", "text/plain"),
        department="IT",
    )
    assert response.status_code == 202
    body_ = response.json()
    assert body_["rejected"] == []
    assert [d["status"] for d in body_["documents"]] == ["queued", "queued"]
    assert all(d["chunk_count"] == 0 and d["creation_date"] is None for d in body_["documents"])
    assert "stored_path" not in response.text, "server file names are never exposed"

    listing = client.get("/api/documents", headers=user).json()
    assert {d["status"] for d in listing["items"]} == {"indexed"}
    assert all(
        d["chunk_count"] >= 1 and d["department"] == "IT" and d["creation_date"]
        for d in listing["items"]
    )
    assert listing["summary"] == {
        "total_files": 2,
        "total_chunks": sum(d["chunk_count"] for d in listing["items"]),
        "by_status": {"indexed": 2},
    }
    assert listing["total"] == 2


def test_a_docx_and_a_pdf_are_accepted_and_indexed(client: TestClient, user: Headers) -> None:
    pdf = make_pdf(
        [[("h", "Purpose"), ("p", f"Alpha {body(1, 40)}")], [("p", f"Bravo {body(2, 40)}")]]
    )
    response = upload(
        client,
        user,
        ("leave.docx", docx_bytes(), "application/octet-stream"),
        ("spec.pdf", pdf, "application/pdf"),
    )
    assert response.status_code == 202
    items = {d["name"]: d for d in client.get("/api/documents", headers=user).json()["items"]}
    assert items["leave.docx"]["status"] == "indexed" and items["spec.pdf"]["status"] == "indexed"
    assert items["spec.pdf"]["mime_type"] == "application/pdf"


def test_a_scanned_pdf_is_accepted_then_fails_with_a_clear_message(
    client: TestClient, user: Headers
) -> None:
    document = one(client, user, "scan.pdf", make_pdf(["image", "image"]), "application/pdf")
    detail = client.get(f"/api/documents/{document['id']}", headers=user).json()
    assert (
        detail["status"] == "failed" and "scanned PDF, OCR not supported" in detail["error_message"]
    )
    assert detail["chunk_count"] == 0


def test_refused_files_are_reported_without_blocking_the_good_ones(
    client: TestClient, user: Headers
) -> None:
    response = upload(
        client,
        user,
        ("good.md", MD, "text/markdown"),
        ("malware.exe", b"MZ\x90\x00", "application/octet-stream"),
        ("fake.pdf", b"just some text, not a pdf", "application/pdf"),
    )
    assert response.status_code == 202
    body_ = response.json()
    assert [d["name"] for d in body_["documents"]] == ["good.md"]
    reasons = {r["filename"]: r["reason"] for r in body_["rejected"]}
    assert "Unsupported file type '.exe'" in reasons["malware.exe"]
    assert reasons["fake.pdf"] == "The file is not a valid PDF."


def test_an_upload_where_nothing_is_acceptable_is_a_400_with_reasons(
    client: TestClient, user: Headers
) -> None:
    response = upload(
        client, user, ("a.exe", b"MZ", "application/octet-stream"), ("empty.txt", b"", "text/plain")
    )
    assert response.status_code == 400
    detail = response.json()["detail"]
    assert detail["message"] == "No file was accepted."
    assert {r["filename"] for r in detail["rejected"]} == {"a.exe", "empty.txt"}
    assert any("empty" in r["reason"] for r in detail["rejected"])


def test_the_size_limit_is_enforced(
    client: TestClient, user: Headers, monkeypatch: pytest.MonkeyPatch, upload_dir: Path
) -> None:
    monkeypatch.setattr(get_settings(), "max_upload_mb", 1)
    response = upload(
        client,
        user,
        ("big.txt", b"word " * 250_000, "text/plain"),
        ("ok.txt", b"A small note about retries.", "text/plain"),
    )
    body_ = response.json()
    assert [d["name"] for d in body_["documents"]] == ["ok.txt"]
    assert "1 MB limit" in body_["rejected"][0]["reason"]
    assert len(list(upload_dir.iterdir())) == 1, "the oversized file must not be left on disk"


@pytest.mark.parametrize(
    ("filename", "shown"),
    [
        ("../../../etc/passwd.txt", "passwd.txt"),
        ("..\\..\\evil.md", "evil.md"),
        ("C:\\Windows\\win.ini.txt", "win.ini.txt"),
        ("a/b/c/deep.txt", "deep.txt"),
    ],
)
def test_user_supplied_paths_are_never_used_to_store_files(
    client: TestClient, user: Headers, upload_dir: Path, tmp_path: Path, filename: str, shown: str
) -> None:
    document = one(client, user, filename, b"Some harmless text content here.", "text/plain")
    assert document["name"] == shown
    stored = [p.name for p in upload_dir.iterdir()]
    assert len(stored) == 1 and re.fullmatch(r"[0-9a-f]{32}\.(txt|md)", stored[0])
    assert {p.name for p in tmp_path.iterdir()} == {"uploads"}, (
        "nothing may be written outside the upload directory"
    )


def test_the_declared_content_type_is_checked_but_not_trusted(
    client: TestClient, user: Headers
) -> None:
    pdf = make_pdf([[("p", body(1, 30))]])
    bad = upload(client, user, ("report.pdf", pdf, "image/png"))
    assert (
        bad.status_code == 400 and "content type" in bad.json()["detail"]["rejected"][0]["reason"]
    )
    assert upload(client, user, ("report.pdf", pdf, "application/octet-stream")).status_code == 202
    assert upload(client, user, ("report2.pdf", pdf, "")).status_code == 202


@pytest.mark.parametrize(
    ("name", "content", "reason"),
    [
        ("fake.docx", b"plain text pretending to be a docx", "not a valid DOCX"),
        ("zip.docx", b"PK\x03\x04" + b"\x00" * 40, "not a valid DOCX"),
        ("fake.pdf", b"PK\x03\x04 zip bytes", "not a valid PDF"),
        ("bin.txt", b"abc\x00\x01\x02def" * 20, "does not look like a text document"),
    ],
)
def test_file_contents_must_match_the_extension(
    client: TestClient, user: Headers, name: str, content: bytes, reason: str
) -> None:
    response = upload(client, user, (name, content, "application/octet-stream"))
    assert (
        response.status_code == 400 and reason in response.json()["detail"]["rejected"][0]["reason"]
    )


def test_a_non_admin_cannot_upload_restricted_documents(
    client: TestClient, user: Headers, upload_dir: Path, db: Session
) -> None:
    response = upload(client, user, ("secret.md", MD, "text/markdown"), access_level="Restricted")
    assert response.status_code == 403 and "administrators" in response.json()["detail"]
    assert not upload_dir.exists() or list(upload_dir.iterdir()) == []
    assert db.scalar(select(func.count()).select_from(Document)) == 0


def test_an_admin_can_upload_restricted_documents_and_a_user_cannot_see_them(
    client: TestClient, admin: Headers, user: Headers
) -> None:
    restricted = one(client, admin, "board.md", MD, access_level="Restricted")
    assert restricted["access_level"] == "Restricted"
    assert [d["id"] for d in client.get("/api/documents", headers=admin).json()["items"]] == [
        restricted["id"]
    ]
    listing = client.get("/api/documents", headers=user).json()
    assert listing["items"] == [] and listing["summary"]["total_files"] == 0


@pytest.mark.parametrize(
    "data",
    [
        {"access_level": "secret"},
        {"creation_date": "yesterday"},
        {"creation_date": "2026-13-45"},
        {"department": "x" * 101},
    ],
)
def test_invalid_metadata_is_a_422_and_stores_nothing(
    client: TestClient, user: Headers, db: Session, data: dict[str, str]
) -> None:
    response = upload(client, user, ("a.md", MD, "text/markdown"), **data)
    assert response.status_code == 422
    assert db.scalar(select(func.count()).select_from(Document)) == 0


def test_explicit_metadata_is_canonicalized_and_wins(client: TestClient, user: Headers) -> None:
    document = one(
        client,
        user,
        "HR_x_2020.md",
        MD,
        department=" legal ",
        document_type="contract",
        creation_date="2024-02-29",
        access_level="public",
    )
    detail = client.get(f"/api/documents/{document['id']}", headers=user).json()
    assert (
        detail["department"],
        detail["document_type"],
        detail["creation_date"],
        detail["access_level"],
    ) == ("Legal", "Contract", "2024-02-29", "Public")


def test_a_missing_files_field_is_a_422(client: TestClient, user: Headers) -> None:
    assert client.post("/api/documents", headers=user, data={"department": "IT"}).status_code == 422


def test_every_route_requires_authentication(client: TestClient) -> None:
    any_id = "00000000-0000-0000-0000-000000000001"
    assert (
        client.post("/api/documents", files=[("files", ("a.md", MD, "text/markdown"))]).status_code
        == 401
    )
    assert client.get("/api/documents").status_code == 401
    assert client.get(f"/api/documents/{any_id}").status_code == 401
    assert client.get(f"/api/documents/{any_id}/file").status_code == 401
    assert client.delete(f"/api/documents/{any_id}").status_code == 401


# --- Listing and filtering ---------------------------------------------------------------------------------


def test_list_filters_search_and_paging(client: TestClient, user: Headers) -> None:
    one(client, user, "HR_Leave_Policy.md", MD, department="HR")
    one(client, user, "IT_Network_Spec.md", MD, department="IT")
    one(client, user, "IT_Backup_Policy.md", MD, department="IT")

    def names(**params: str) -> list[str]:
        items = client.get("/api/documents", headers=user, params=params).json()["items"]
        return sorted(d["name"] for d in items)

    assert names(department="it") == ["IT_Backup_Policy.md", "IT_Network_Spec.md"]
    assert names(department="HR") == ["HR_Leave_Policy.md"]
    assert names(q="policy") == ["HR_Leave_Policy.md", "IT_Backup_Policy.md"]
    assert names(q="100%") == [] and names(q="%") == [] and names(q="L_ave") == [], (
        "LIKE wildcards in the search text are literal"
    )
    assert names(q="_") != [], "a literal underscore does match the underscores in these names"
    assert names(status="indexed") == sorted(names()) and names(status="failed") == []
    page = client.get("/api/documents", headers=user, params={"limit": 2, "offset": 0}).json()
    assert len(page["items"]) == 2 and page["total"] == 3 and page["summary"]["total_files"] == 3
    rest = client.get("/api/documents", headers=user, params={"limit": 2, "offset": 2}).json()
    assert len(rest["items"]) == 1
    assert client.get("/api/documents", headers=user, params={"status": "bogus"}).status_code == 422
    assert client.get("/api/documents", headers=user, params={"limit": 0}).status_code == 422


def test_the_library_summary_ignores_filters_and_counts_failures(
    client: TestClient, user: Headers
) -> None:
    one(client, user, "ok.md", MD)
    one(client, user, "scan.pdf", make_pdf(["image"]), "application/pdf")
    filtered = client.get("/api/documents", headers=user, params={"status": "failed"}).json()
    assert len(filtered["items"]) == 1
    assert (
        filtered["summary"]["by_status"] == {"indexed": 1, "failed": 1}
        and filtered["summary"]["total_files"] == 2
    )


# --- Visibility -----------------------------------------------------------------------------------------------


def test_documents_above_your_level_do_not_exist_for_you(
    client: TestClient, admin: Headers, user: Headers
) -> None:
    secret = one(client, admin, "board.md", MD, access_level="Restricted")
    for path in (f"/api/documents/{secret['id']}", f"/api/documents/{secret['id']}/file"):
        assert client.get(path, headers=user).status_code == 404, (
            "404, never 403: the document's existence is not revealed"
        )
        assert client.get(path, headers=admin).status_code == 200
    assert client.delete(f"/api/documents/{secret['id']}", headers=user).status_code == 404
    assert client.get(f"/api/documents/{secret['id']}", headers=admin).json()["name"] == "board.md"


def test_unknown_and_malformed_ids(client: TestClient, user: Headers) -> None:
    assert (
        client.get("/api/documents/00000000-0000-0000-0000-000000000009", headers=user).status_code
        == 404
    )
    assert client.get("/api/documents/not-a-uuid", headers=user).status_code == 422


# --- The original file -------------------------------------------------------------------------------------------


def test_a_pdf_is_served_inline_with_safe_headers(client: TestClient, user: Headers) -> None:
    pdf = make_pdf([[("p", body(1, 40))]])
    document = one(client, user, "Spec Sheet é.pdf", pdf, "application/pdf")
    response = client.get(f"/api/documents/{document['id']}/file", headers=user)
    assert response.status_code == 200 and response.content == pdf
    assert response.headers["content-type"] == "application/pdf"
    assert response.headers["content-disposition"].startswith("inline")
    assert (
        response.headers["x-content-type-options"] == "nosniff"
        and "no-store" in response.headers["cache-control"]
    )
    forced = client.get(
        f"/api/documents/{document['id']}/file", headers=user, params={"download": "true"}
    )
    assert forced.headers["content-disposition"].startswith("attachment")


@pytest.mark.parametrize(
    ("name", "content"),
    [
        ("page.html", b"<html><body><p>Hello <script>alert(1)</script></p></body></html>"),
        ("notes.md", MD),
        ("plain.txt", b"Some plain text here."),
    ],
)
def test_everything_except_pdf_is_a_download_and_never_rendered_inline(
    client: TestClient, user: Headers, name: str, content: bytes
) -> None:
    """An uploaded HTML page must not run script on our origin, next to the session token."""
    document = one(client, user, name, content, "text/html")
    response = client.get(f"/api/documents/{document['id']}/file", headers=user)
    assert response.status_code == 200 and response.content == content
    assert response.headers["content-disposition"].startswith("attachment")
    assert "sandbox" in response.headers["content-security-policy"]


def test_a_missing_original_is_a_clean_404(
    client: TestClient, user: Headers, upload_dir: Path
) -> None:
    document = one(client, user)
    for path in upload_dir.iterdir():
        path.unlink()
    response = client.get(f"/api/documents/{document['id']}/file", headers=user)
    assert response.status_code == 404 and "no longer available" in response.json()["detail"]


def test_the_media_type_comes_from_our_validation_not_the_client(
    client: TestClient, user: Headers
) -> None:
    document = one(client, user, "notes.txt", b"Some plain text here.", "application/x-evil")
    assert document["mime_type"] == "text/plain"


# --- Deleting --------------------------------------------------------------------------------------------------------


def assert_gone(db: Session, upload_dir: Path, document_id: str) -> None:
    db.expire_all()
    assert db.get(Document, document_id) is None
    assert (
        db.scalar(select(func.count()).select_from(Chunk).where(Chunk.document_id == document_id))
        == 0
    )
    assert vectorstore.ids_for_document(document_id) == []
    assert not [c for c in bm25.get_index().chunk_ids() if c.startswith(document_id)]
    assert list(upload_dir.iterdir()) == []


def test_the_owner_can_delete_and_nothing_remains_anywhere(
    client: TestClient,
    user: Headers,
    db: Session,
    upload_dir: Path,
    fake_embeddings: FakeEmbeddings,
) -> None:
    document = one(client, user, "policy.md", MD)
    assert vectorstore.count() >= 1 and len(bm25.get_index()) >= 1
    assert client.delete(f"/api/documents/{document['id']}", headers=user).status_code == 204
    assert_gone(db, upload_dir, document["id"])
    assert client.get(f"/api/documents/{document['id']}", headers=user).status_code == 404
    assert client.delete(f"/api/documents/{document['id']}", headers=user).status_code == 404
    # its content can no longer be retrieved by either channel
    assert vectorstore.query(fake_embeddings.embed_query("retry limit API timeouts"), 5) == []
    assert bm25.get_index().search("retry limit timeouts", 5) == []


def test_only_the_owner_or_an_admin_may_delete(
    client: TestClient, admin: Headers, user: Headers, other: Headers, db: Session, upload_dir: Path
) -> None:
    mine = one(client, user, "mine.md", MD)
    forbidden = client.delete(f"/api/documents/{mine['id']}", headers=other)
    assert (
        forbidden.status_code == 403
        and "administrator or the uploader" in forbidden.json()["detail"]
    )
    assert client.get(f"/api/documents/{mine['id']}", headers=other).status_code == 200, (
        "others can still read it"
    )
    assert client.delete(f"/api/documents/{mine['id']}", headers=admin).status_code == 204
    assert_gone(db, upload_dir, mine["id"])


def test_deleting_one_document_leaves_the_others_searchable(
    client: TestClient, user: Headers, fake_embeddings: FakeEmbeddings
) -> None:
    keep = one(
        client, user, "keep.md", b"# Keep\n\nThe quokka migration schedule is published annually.\n"
    )
    drop = one(client, user, "drop.md", MD)
    assert client.delete(f"/api/documents/{drop['id']}", headers=user).status_code == 204
    assert {m.chunk_id.split("_")[0] for m in bm25.get_index().search("quokka migration", 5)} == {
        keep["id"]
    }
    assert {
        h.metadata["document_id"]
        for h in vectorstore.query(fake_embeddings.embed_query("quokka"), 5)
    } == {keep["id"]}


def test_a_document_being_indexed_cannot_be_deleted(
    client: TestClient, user: Headers, db: Session
) -> None:
    document = one(client, user)
    lock = ingestion.document_lock(db.scalars(select(Document.id)).one())
    assert lock.acquire(blocking=False)
    try:
        response = client.delete(f"/api/documents/{document['id']}", headers=user)
    finally:
        lock.release()
    assert response.status_code == 409 and "being indexed" in response.json()["detail"]
    assert client.get(f"/api/documents/{document['id']}", headers=user).status_code == 200


def test_a_failed_deletion_parks_the_document_and_can_be_retried(
    client: TestClient,
    user: Headers,
    db: Session,
    upload_dir: Path,
    monkeypatch: pytest.MonkeyPatch,
    fake_embeddings: FakeEmbeddings,
) -> None:
    document = one(client, user)
    real_delete = vectorstore.delete_document

    def down(_: str) -> None:
        raise vectorstore.VectorStoreError("ChromaDB delete failed: ConnectError")

    monkeypatch.setattr(vectorstore, "delete_document", down)
    response = client.delete(f"/api/documents/{document['id']}", headers=user)
    assert response.status_code == 503 and "Delete the document again" in response.json()["detail"]
    assert "ConnectError" not in response.text

    parked = client.get(f"/api/documents/{document['id']}", headers=user).json()
    assert parked["status"] == "failed" and "Deletion did not finish" in parked["error_message"]
    assert bm25.get_index().search("retry limit timeouts", 5) == [], (
        "never half-searchable, even mid-failure"
    )

    monkeypatch.setattr(vectorstore, "delete_document", real_delete)
    assert client.delete(f"/api/documents/{document['id']}", headers=user).status_code == 204
    assert_gone(db, upload_dir, document["id"])


# --- Upload validation added by the audit ----------------------------------------------------------------


@pytest.mark.parametrize("bad", ["1985-06-01", "1989-12-31", "2099-01-01"])
def test_an_implausible_creation_date_is_refused_instead_of_silently_replaced(
    client: TestClient, admin: Headers, bad: str
) -> None:
    response = upload(client, admin, ("a.md", MD, "text/markdown"), creation_date=bad)
    assert response.status_code == 422
    assert "creation_date" in response.json()["detail"] and "1990" in response.json()["detail"]
    assert (
        upload(client, admin, ("a.md", MD, "text/markdown"), creation_date="2026-01-15").status_code
        == 202
    )


def test_a_request_with_too_many_files_is_refused_whole(
    client: TestClient, admin: Headers, monkeypatch: pytest.MonkeyPatch, upload_dir: Path
) -> None:
    monkeypatch.setattr(get_settings(), "max_files_per_request", 2)
    files = [(f"d{n}.md", MD, "text/markdown") for n in range(3)]
    response = upload(client, admin, *files)
    assert response.status_code == 413 and "At most 2 files" in response.json()["detail"]
    assert list(upload_dir.glob("*")) == []  # nothing was stored
    assert upload(client, admin, *files[:2]).status_code == 202


def test_a_failure_while_storing_a_later_file_leaves_nothing_on_disk(
    db: Session, upload_dir: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A full disk on the third file must not strand the first two on disk with no row."""
    from app.core.metadata import ExplicitMetadata
    from app.db.models import User
    from app.services import document_service
    from app.services.document_service import IncomingFile

    owner = User(email="u@example.com", password_hash="x", role="user")
    db.add(owner)
    db.commit()
    real, calls = document_service._store_file, {"n": 0}

    def flaky(file: IncomingFile, extension: str, max_bytes: int) -> tuple[str, int]:
        calls["n"] += 1
        if calls["n"] == 3:
            raise OSError("no space left on device")
        return real(file, extension, max_bytes)

    monkeypatch.setattr(document_service, "_store_file", flaky)
    files = [IncomingFile(f"d{n}.md", None, io.BytesIO(MD), len(MD)) for n in range(4)]
    with pytest.raises(OSError):
        document_service.create_documents(
            db, owner, files, ExplicitMetadata(), submit=lambda _id: None
        )
    assert list(upload_dir.glob("*")) == []
    assert db.scalar(select(func.count()).select_from(Document)) == 0
